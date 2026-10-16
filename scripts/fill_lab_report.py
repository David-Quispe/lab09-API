"""Complete the guided sections of the supplied Lab 09 Word template."""

from __future__ import annotations

import argparse
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from PIL import Image


ANSWERS = {
    1: (
        "Además de /posts y /posts/{id}, JSONPlaceholder Mock ofrece /comments, "
        "/companies, /users, /roles y /todos, cada uno con una ruta para listar "
        "y otra para consultar un elemento por ID. También incluye POST /login y GET / "
        "para los metadatos. Fuente: https://app.beeceptor.com/virtual-sandbox/json-placeholder"
    ),
    5: (
        "@SerializedName indica a Gson el nombre exacto de la propiedad JSON que debe asociar "
        "con un atributo del modelo Kotlin. Por ejemplo, comment_count se convierte en "
        "commentCount; así se conserva la convención de nombres de Kotlin sin perder el mapeo."
    ),
    7: (
        "@GET indica a Retrofit que el método hace una petición HTTP GET a la ruta indicada. "
        "@GET(\"posts\") consulta la lista y @GET(\"posts/{id}\") consulta un post; "
        "@Path(\"id\") sustituye el segmento variable por el ID recibido."
    ),
    8: (
        "suspend declara una función que puede suspenderse dentro de una corrutina mientras "
        "espera la respuesta de la red, sin bloquear el hilo principal. Retrofit adapta la "
        "llamada HTTP a esa función; suspend por sí solo no crea un hilo nuevo."
    ),
    10: (
        "rememberNavController() crea y recuerda un NavHostController para la composición. "
        "Este controlador conserva el estado de navegación durante recomposiciones y permite "
        "navegar entre inicio, posts y detalle mediante NavHost y navController.navigate()."
    ),
    14: (
        "Contenido recibe el espacio interior (PaddingValues) de Scaffold y lo aplica al Box "
        "para que las pantallas no queden debajo de las barras. Dentro, NavHost define inicio "
        "como ruta inicial, posts para la lista y postsVer/{id} para el detalle. El argumento "
        "id se interpreta como entero y se pasa a ScreenPost para consultar el post elegido."
    ),
    16: (
        "LaunchedEffect ejecuta una corrutina vinculada a la composición. Se inicia al entrar "
        "en pantalla, se reinicia si cambia su clave y se cancela al salir. Aquí dispara "
        "loadPosts() en la lista y loadPost(id) en el detalle."
    ),
    17: (
        "SnapshotStateList es una lista mutable observable por Compose: al añadir, quitar o "
        "cambiar elementos, las partes de la interfaz que la leen se recomponen. En esta "
        "versión MVVM se usa una lista inmutable dentro de PostsUiState, expuesta mediante "
        "StateFlow, para que el ViewModel controle el estado."
    ),
    18: (
        "LazyColumn muestra una lista vertical desplazable y compone principalmente los "
        "elementos visibles. Se usa para representar los posts con items(state.posts), "
        "identificando cada fila por su id."
    ),
    19: (
        "Al tocar una fila, ScreenPosts navega a postsVer/{id} con el ID del elemento. "
        "NavHost extrae el entero y abre ScreenPost. LaunchedEffect(id) pide al ViewModel "
        "cargarlo; el repositorio llama GET posts/{id} y el nuevo PostsUiState hace que "
        "Compose muestre su ID, usuario, título y cuerpo."
    ),
    20: (
        "Sí. La capa de datos contiene PostModel, PostApiService y PostRepository; el "
        "repositorio configura Retrofit y ofrece getPosts()/getPost(id). PostsViewModel "
        "consulta esa capa y publica PostsUiState mediante StateFlow. MainActivity y "
        "moduloPosts.kt observan el estado y dibujan las pantallas; las acciones de la UI "
        "llaman al ViewModel."
    ),
}


def insert_paragraph_after(document: Document, anchor) -> Paragraph:
    paragraph = document.add_paragraph()
    anchor.addnext(paragraph._p)
    return paragraph


def add_evidence(document: Document, anchor, label: str, image_path: Path, width: float, crop=None):
    caption = insert_paragraph_after(document, anchor)
    caption.add_run(label).bold = True
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(3)
    caption.paragraph_format.keep_with_next = True
    caption.style = "Normal"
    for run in caption.runs:
        run.font.size = Pt(9)

    figure = insert_paragraph_after(document, caption._p)
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.space_after = Pt(7)
    figure.paragraph_format.keep_together = True

    if crop:
        with TemporaryDirectory(dir=image_path.parent.parent) as temporary_directory:
            final_path = Path(temporary_directory) / image_path.name
            with Image.open(image_path) as source:
                source.crop(crop).save(final_path)
            figure.add_run().add_picture(str(final_path), width=Inches(width))
    else:
        figure.add_run().add_picture(str(image_path), width=Inches(width))
    return figure._p


def fill_answer(document: Document, table_number: int, text: str):
    cell = document.tables[table_number].cell(0, 0)
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.size = Pt(10)


def add_hyperlink(paragraph: Paragraph, label: str, url: str):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    document = Document(args.source)
    original_paragraphs = list(document.paragraphs)
    workspace = Path(__file__).resolve().parents[1]
    evidence = workspace / "evidencias"

    for table_number, answer in ANSWERS.items():
        fill_answer(document, table_number, answer)

    # Keep the instructor's code examples and place proof of the actual files beside them.
    add_evidence(document, original_paragraphs[40]._p,
                 "Evidencia 1. Sitio JSONPlaceholder Mock consultado.",
                 evidence / "01-api.jpg", 2.4)
    add_evidence(document, document.tables[2]._element,
                 "Evidencia 2. Permiso de Internet en AndroidManifest.xml.",
                 evidence / "02-permiso-internet.png", 6.1, (265, 44, 1140, 420))
    add_evidence(document, document.tables[3]._element,
                 "Evidencia 3. Versiones y bibliotecas de Navigation, ViewModel y Retrofit.",
                 evidence / "02-dependencias.png", 6.1, (265, 44, 1340, 870))
    add_evidence(document, document.tables[4]._element,
                 "Evidencia 4. PostModel implementado con @SerializedName.",
                 evidence / "03-postmodel.png", 6.1, (265, 44, 1120, 510))
    add_evidence(document, document.tables[6]._element,
                 "Evidencia 5. PostApiService con los endpoints GET.",
                 evidence / "04-servicio-api.png", 6.1, (265, 44, 1150, 465))
    add_evidence(document, document.tables[9]._element,
                 "Evidencia 6. Retrofit, URL base, conversor Gson e instancia del servicio.",
                 evidence / "05-retrofit.png", 6.1, (265, 44, 1160, 560))
    add_evidence(document, document.tables[11]._element,
                 "Evidencia 7. ProgPrincipal9 y la estructura Scaffold.",
                 evidence / "05-scaffold.png", 6.1, (265, 44, 1230, 750))
    add_evidence(document, document.tables[12]._element,
                 "Evidencia 8. Contenido y las rutas de navegación.",
                 evidence / "05-contenido.png", 6.1, (265, 44, 1250, 930))
    anchor = add_evidence(document, document.tables[15]._element,
                 "Evidencia 9. LazyColumn, LaunchedEffect y selección de un post.",
                 evidence / "06-lista-codigo.png", 6.1, (265, 44, 1280, 850))
    add_evidence(document, anchor,
                 "Evidencia 10. ScreenPost y visualización del detalle.",
                 evidence / "06-detalle-codigo.png", 6.1, (265, 44, 1280, 930))
    add_evidence(document, document.tables[20]._element,
                 "Evidencia 11. PostsViewModel y estado de interfaz MVVM.",
                 evidence / "06-mvvm.png", 6.1, (265, 44, 1250, 930))

    anchor = original_paragraphs[93]._p
    for label, name in [
        ("Evidencia 12. Pantalla de inicio ejecutada en el emulador.", "08-inicio.png"),
        ("Evidencia 13. Lista de posts recibida desde la API.", "08-posts.png"),
        ("Evidencia 14. Detalle del post 1 recibido desde la API.", "08-detalle.png"),
    ]:
        anchor = add_evidence(document, anchor, label, evidence / name, 2.5)

    video = original_paragraphs[94]
    video.clear()
    video.add_run("Video de funcionamiento (MP4): ").bold = True
    add_hyperlink(
        video,
        "ver grabación en GitHub",
        "https://github.com/David-Quispe/lab09-API/blob/main/evidencias/09-funcionamiento.mp4",
    )
    video.paragraph_format.space_after = Pt(6)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)


if __name__ == "__main__":
    main()
